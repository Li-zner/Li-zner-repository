"""查看最新评测汇总"""
import openpyxl

wb = openpyxl.load_workbook("tests/evaluation_report.xlsx")
ws = wb.active

total = ws.max_row - 1
print(f"总用例数: {total}")

# 按分类统计
categories = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    cid = str(row[0] or "")
    score = row[3]
    if cid.startswith("CC"):
        cat = "民法典"
    elif cid.startswith("TR"):
        cat = "旅游"
    elif cid.startswith("RB"):
        cat = "鲁棒性"
    elif cid.startswith("AT"):
        cat = "攻击测试"
    else:
        cat = "其他"
    
    if cat not in categories:
        categories[cat] = {"count": 0, "scored": 0, "total_score": 0}
    categories[cat]["count"] += 1
    if score not in (None, "-"):
        try:
            s = float(score)
            categories[cat]["scored"] += 1
            categories[cat]["total_score"] += s
        except (ValueError, TypeError):
            pass

print(f"\n📊 分类统计:")
all_scored = 0
all_total = 0
for cat, data in sorted(categories.items()):
    avg = data["total_score"] / data["scored"] if data["scored"] > 0 else 0
    print(f"  {cat}: {data['count']}条, 评分{data['scored']}条, 平均{avg:.2f}")
    all_scored += data["total_score"]
    all_total += data["scored"]

if all_total > 0:
    print(f"\n  📈 总平均: {all_scored/all_total:.2f}/10 ({total}条)")

wb.close()
