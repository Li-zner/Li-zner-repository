"""查看最新测评结果"""
import openpyxl, yaml

# Load test cases
with open(r'D:\桌面\agent_gateway\tests\test_cases.yaml', 'r', encoding='utf-8') as f:
    raw = yaml.safe_load(f)
all_cases = raw.get('test_cases', raw)
cc_ids = {c['id'] for c in all_cases if c.get('category') == '民法典'}

# Load report
wb = openpyxl.load_workbook(r'D:\桌面\agent_gateway\tests\civil_evaluation_report.xlsx')
ws = wb.active
header = [cell.value for cell in ws[1]]
id_col = 0
score_col = header.index('评分(10分制)')
analysis_col = header.index('扣分说明与综合评价')
resp_col = header.index('Agent完整回复')
err_col = header.index('错误')

cc_rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    cid = row[id_col]
    if cid in cc_ids:
        cc_rows.append((cid, row[score_col], str(row[analysis_col] or ''), str(row[resp_col] or ''), str(row[err_col] or '')))

cc_rows.sort(key=lambda x: x[0])

scores = [s for _, s, _, _, _ in cc_rows]
avg = sum(scores) / len(scores)
high = sum(1 for s in scores if s >= 8)
mid = sum(1 for s in scores if 6 <= s < 8)
low = sum(1 for s in scores if s < 6)

print(f"Total: {len(cc_rows)}")
print(f"Average: {avg:.2f}/10")
print(f"Max: {max(scores)}, Min: {min(scores)}")
print(f"High(>=8): {high}, Mid(6-8): {mid}, Low(<6): {low}")
print()

# Focus on previously problematic cases
targets = {'CC001': '截断', 'CC032': '截断', 'CC035': '截断', 'CC057': '截断', 'CC058': '安全过滤', 'CC060': '编造法条'}
print("=== 之前有问题的用例 ===")
for cid, score, analysis, resp, err in cc_rows:
    if cid in targets:
        print(f"\n--- {cid} ({targets[cid]}): {score}/10 ---")
        print(f"  响应长度: {len(resp)} 字符")
        print(f"  响应开头: {resp[:150]}")
        print(f"  评价摘要: {analysis[:200]}")
        if err and err.strip() and err != '-':
            print(f"  错误: {err}")
print()

# All scores
print("=== 所有分数 ===")
for cid, score, _, _, _ in cc_rows:
    marker = ""
    if cid in targets:
        marker = " ***" + targets[cid]
    elif score < 8:
        marker = " ***低分"
    print(f"  {cid}: {score}/10{marker}")
