"""快速检查民法典测评结果 - 从excel报告读取"""
import openpyxl
import yaml

# Load test cases to get categories
with open(r'D:\桌面\agent_gateway\tests\test_cases.yaml', 'r', encoding='utf-8') as f:
    raw = yaml.safe_load(f)

all_cases = raw.get('test_cases', raw)
if isinstance(all_cases, dict):
    all_cases = list(all_cases.values())

cc_ids = {c['id'] for c in all_cases if c.get('category') == '民法典'}

wb = openpyxl.load_workbook(r'D:\桌面\agent_gateway\tests\evaluation_report.xlsx')
ws = wb.active

header = [cell.value for cell in ws[1]]
id_col = 0
score_col = header.index('评分(10分制)')
analysis_col = header.index('扣分说明与综合评价')

# Collect civil code scores
cc_rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    cid = row[id_col]
    if cid in cc_ids:
        cc_rows.append((cid, row[score_col], row[analysis_col]))

cc_rows.sort(key=lambda x: x[0])

print(f"CC cases scored: {len(cc_rows)}")
if cc_rows:
    scores = [s for _, s, _ in cc_rows]
    avg = sum(scores) / len(scores)
    high = sum(1 for s in scores if s >= 8)
    mid = sum(1 for s in scores if 6 <= s < 8)
    low = sum(1 for s in scores if s < 6)
    print(f"Average: {avg:.2f}/10")
    print(f"Max: {max(scores)}, Min: {min(scores)}")
    print(f"High(>=8): {high}, Mid(6-8): {mid}, Low(<6): {low}")
    print()
    
    low_cases = [x for x in cc_rows if x[1] < 8]
    print(f"Low score cases (<8): {len(low_cases)}")
    for cid, score, analysis in low_cases:
        print(f"\n=== {cid}: {score}/10 ===")
        if analysis:
            print(f"  {analysis[:200]}")
        else:
            print("  (no analysis)")
    print()
    
    # All scores list
    print("All CC scores:")
    for cid, score, _ in cc_rows:
        marker = " ***" if score < 8 else ""
        print(f"  {cid}: {score}/10{marker}")
