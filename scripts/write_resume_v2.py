# -*- coding: utf-8 -*-
"""简历第二轮优化：在 黎忠南简历-优化版.docx 基础上
1) 加教育背景一行（知识库已确认：广西科技大学/本科/土木工程/2021-2025）
2) 按 HR 深挖点加固：可插拔架构写具体（JSON 配置包）、"约六成"标注压测依据、
   源站零端口暴露改写为安全决策、补充"3个月从零到上线+一人独立"、新增开源与自研澄清
仅改文本，保留原段落样式与格式。"""
from docx import Document
from docx.shared import Pt

SRC = r"d:\桌面\agent_gateway\黎忠南简历-优化版.docx"
DST = r"d:\桌面\agent_gateway\黎忠南简历-最终版.docx"

doc = Document(SRC)


def set_para_text(p, new_text):
    """保留段落第一个 run 的格式，其余 run 清空，写入新文本。"""
    if p.runs:
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new_text)


def replace_first(paragraphs, old_frag, new_text):
    """找第一个包含 old_frag 的段落并整体替换文本。返回是否成功。"""
    for p in paragraphs:
        if old_frag in p.text:
            set_para_text(p, new_text)
            return True
    return False


def insert_before(paragraphs, anchor_frag, new_text, size=9.0, bold=None):
    """在第一个包含 anchor_frag 的段落之前插入新段落，指定字号。返回是否成功。"""
    for p in paragraphs:
        if anchor_frag in p.text:
            new_p = p.insert_paragraph_before()
            run = new_p.add_run(new_text)
            run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            return True
    return False


# ============ 1. 教育背景（插在"求职意向"之后、"个人简介"之前） ============
insert_before(
    doc.paragraphs,
    "个人简介",
    "教育背景：广西科技大学 · 本科 · 土木工程（土木工程建造与管理）｜2021-2025 ｜ 曾任校组织运营部副部长、班级心理委员",
    size=9.0,
)

# ============ 2. 核心优势：可插拔架构写具体（防"配置包是什么格式"追问） ============
replace_first(
    doc.paragraphs,
    "引擎与业务解耦的可插拔架构",
    "•架构与复用：引擎与业务解耦的可插拔架构，人设、工具选择、模型策略全部外置为 JSON 配置包；"
    "新增业务场景仅需新增一份配置包、核心代码零改动，具备产品化落地意识。",
)

# ============ 3. 项目经历开头：补"3个月从零到上线、一人独立完成" ============
replace_first(
    doc.paragraphs,
    "面向企业知识服务",
    "面向企业知识服务的 AI 网关，含完整支付体系，已上线公网稳定运行；"
    "3 个月从零到上线、全程一人独立完成。本地压测：800 并发近零错误、极限 1600 并发可用性 99.28%。",
)

# ============ 4. 架构设计条目：业务能力通过 JSON 配置包接入 ============
replace_first(
    doc.paragraphs,
    "业务能力通过配置包接入",
    "•架构设计：独立完成从需求拆解、架构设计到部署上线的全流程；"
    "Agent 内核只负责调度决策与流式输出，业务能力通过 JSON 配置包接入，支撑多场景扩展。",
)

# ============ 5. 性能优化条目："约六成"标注压测依据（防"怎么算的"追问） ============
replace_first(
    doc.paragraphs,
    "三级语义缓存将命中查询成本",
    "•性能优化：三级语义缓存将命中查询成本降至纯调用的约六成（本地压测对比）；"
    "冷热分层记忆 + 断连自动保存；本地压测 800 并发平均响应 1.1s。",
)

# ============ 6. 可观测与运维条目：零端口暴露改写为安全决策 ============
replace_first(
    doc.paragraphs,
    "Cloudflare Tunnel 公网暴露",
    "•可观测与运维：Docker Compose 编排 10+ 服务、4 实例 + Nginx 负载均衡与按 QPS 自动伸缩；"
    "Cloudflare Tunnel 安全隧道公网暴露、源站不开放入站端口、密钥全部外置环境变量；慢查询监控与异常聚合告警。",
)

# ============ 7. 新增"开源与自研"条目（回答用了多少开源/哪些自己写） ============
insert_before(
    doc.paragraphs,
    "技术栈",
    "•开源与自研：基于 FastAPI / PostgreSQL / Redis / LangChain 等开源组件搭建；"
    "核心编排、支付并发安全（分布式锁 / 乐观锁 / WAL 流水）、三级语义缓存、限流熔断等关键逻辑全部自行实现。",
    size=9.5,
)

doc.save(DST)
print("saved:", DST)

# ============ 校验 ============
check = Document(DST)
print("---- 段落数:", len(check.paragraphs))
for p in check.paragraphs:
    t = p.text.strip()
    if t:
        print(" |", t[:100])
