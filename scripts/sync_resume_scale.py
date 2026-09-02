# -*- coding: utf-8 -*-
"""同步 3 处口径：
1) 简历 docx：可观测条目 "4 实例 + Nginx 负载均衡与按 QPS 自动伸缩" -> "Nginx 负载均衡 + 按 QPS 自动伸缩（2-4 实例）"
2) prompts/me.json：加 GitHub 个人主页；项目数字行同步 2-4 实例口径
me.json 保持 utf-8 无 BOM（json.load 用 utf-8-sig 兼容，写入用 utf-8）"""
import json
from docx import Document

# ========== 1. 简历 docx ==========
DOCX = r"d:\桌面\agent_gateway\黎忠南简历-最终版.docx"
NEW_OPS = (
    "•可观测与运维：Docker Compose 编排 10+ 服务、Nginx 负载均衡 + 按 QPS 自动伸缩（2-4 实例）；"
    "Cloudflare Tunnel 安全隧道公网暴露、源站不开放入站端口、密钥全部外置环境变量；慢查询监控与异常聚合告警。"
)
doc = Document(DOCX)
found = False
for p in doc.paragraphs:
    if "4 实例 + Nginx 负载均衡与按 QPS 自动伸缩" in p.text:
        if p.runs:
            p.runs[0].text = NEW_OPS
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(NEW_OPS)
        found = True
        break
if not found:
    raise SystemExit("docx 未找到可观测条目")
doc.save(DOCX)
print("docx OK")

# ========== 2. prompts/me.json ==========
ME = r"d:\桌面\agent_gateway\prompts\me.json"
with open(ME, "r", encoding="utf-8-sig") as f:
    data = json.load(f)
prompt = data["prompt"]

old_gh = "源码仓库：github.com/Li-zner/Li-zner-repository"
new_gh = "GitHub：github.com/Li-zner ｜ 源码仓库：github.com/Li-zner/Li-zner-repository"
assert old_gh in prompt, "me.json 未找到源码仓库行"
prompt = prompt.replace(old_gh, new_gh)

old_sc = "4 实例 + Nginx 按 QPS 自动伸缩"
new_sc = "Nginx 按 QPS 自动伸缩（2-4 实例）"
assert old_sc in prompt, "me.json 未找到 4 实例行"
prompt = prompt.replace(old_sc, new_sc)

data["prompt"] = prompt
with open(ME, "w", encoding="utf-8", newline="\n") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
    f.write("\n")
print("me.json OK")
print("  GitHub:", new_gh in prompt)
print("  scale :", new_sc in prompt)
