"""检查RAGAS报告"""
import openpyxl

wb = openpyxl.load_workbook("tests/ragas_report_civil_code.xlsx")
ws = wb.active
headers = [cell.value for cell in ws[1]]
print("列:", headers)

total = ws.max_row - 1  # 减去表头
print(f"总数据行: {total}")

# 汇总统计
ar_sum = ac_sum = fa_sum = 0
count = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    ar = row[3] if len(row) > 3 and isinstance(row[3], (int, float)) else 0
    ac = row[4] if len(row) > 4 and isinstance(row[4], (int, float)) else 0
    fa = row[5] if len(row) > 5 and isinstance(row[5], (int, float)) else 0
    ar_sum += ar
    ac_sum += ac
    fa_sum += fa
    count += 1

if count > 0:
    print(f"\n📊 RAGAS 汇总:")
    print(f"  Answer Relevancy:    {ar_sum/count:.4f}")
    print(f"  Answer Correctness:  {ac_sum/count:.4f}")
    print(f"  Faithfulness:        {fa_sum/count:.4f}")
    print(f"  RAGAS综合:           {(ar_sum/count + ac_sum/count + fa_sum/count)/3:.4f}")

wb.close()
