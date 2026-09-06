# -*- coding: utf-8 -*-
"""简历文本大优化：修改 黎忠南简历.docx 为 黎忠南简历-优化版.docx
仅改文本，保留原有段落样式与格式。"""
from docx import Document

SRC = r"d:\桌面\agent_gateway\黎忠南简历.docx"
DST = r"d:\桌面\agent_gateway\黎忠南简历-优化版.docx"

doc = Document(SRC)


def set_para_text(p, new_text):
    """保留段落第一个 run 的格式，其余 run 清空，写入新文本。"""
    if p.runs:
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new_text)


def replace_first(paragraphs, old_frag, new_text):
    """在段落列表中找第一个包含 old_frag 的段落并替换。返回是否成功。"""
    for p in paragraphs:
        if old_frag in p.text:
            set_para_text(p, new_text)
            return True
    return False


# ============ 顶部个人信息区（补姓名、定薪资） ============
replace_first(
    doc.paragraphs,
    "电话",
    "黎忠南  |  电话 13357308241",
)
replace_first(
    doc.paragraphs,
    "期望薪资",
    "求职意向：AI 应用开发 | 期望薪资：10-12K | 期望城市：广州、深圳",
)

# ============ 个人简介 ============
replace_first(
    doc.paragraphs,
    "独立设计并交付企业级",
    "独立设计并交付一套 Multi-Agent 网关（FastAPI + PostgreSQL + Redis + Docker），"
    "一套引擎承载旅游规划、民法典咨询、求职助手三个业务场景，已上线公网稳定运行。"
    "覆盖 AI 应用层完整链路：RAG 检索、Agent 编排、流式输出、语义缓存；"
    "同时具备工程落地能力：并发安全、可观测性、CI/CD、密钥治理与线上排障。"
    "具备从需求拆解、架构设计到生产运维的完整闭环经验。",
)

# ============ 核心优势 ============
replace_first(
    doc.paragraphs,
    "架构与产品化",
    "•架构与复用：引擎与业务解耦的可插拔架构，人设、工具、模型策略全部外置为配置包；"
    "三个业务场景仅通过新增配置接入、核心代码零改动，具备产品化落地意识。",
)
replace_first(
    doc.paragraphs,
    "工程规范",
    "•工程规范：Git 密钥治理（清除 40+ 文件硬编码密钥）、CI/CD 滚动发布与失败回滚、"
    "数据库版本化迁移、SLO 告警、备份与恢复演练，按生产交付标准落地。",
)
replace_first(
    doc.paragraphs,
    "性能与稳定性",
    "•性能与稳定性：分布式限流、熔断降级、三级语义缓存、冷热分层记忆；"
    "本地 Locust 压测 800 并发近零错误、极限 1600 并发可用性 99.28%。",
)
replace_first(
    doc.paragraphs,
    "可观测与排障",
    "•可观测与排障：OpenTelemetry 全链路追踪 + Prometheus/Grafana/Loki 日志关联，"
    "线上问题分钟级定位；曾独立定位并修复配置漂移导致的线上全链路故障。",
)

# ============ 项目经历 ============
replace_first(
    doc.paragraphs,
    "面向企业知识服务",
    "面向企业知识服务的 AI 网关，含完整支付体系，已上线公网稳定运行。"
    "本地压测：800 并发近零错误、极限 1600 并发可用性 99.28%。",
)
replace_first(
    doc.paragraphs,
    "独立完成从需求拆解",
    "•架构设计：独立完成从需求拆解、架构设计到部署上线的全流程；"
    "Agent 内核只负责调度决策与流式输出，业务能力通过配置包接入，支撑多场景扩展。",
)
replace_first(
    doc.paragraphs,
    "三级语义缓存将命中查询成本",
    "•性能优化：三级语义缓存将命中查询成本降至纯调用的约六成；"
    "冷热分层记忆 + 断连自动保存；本地压测 800 并发平均响应 1.1s。",
)

# ============ 技术栈（熟练只留核心，其余降级） ============
replace_first(
    doc.paragraphs,
    "熟练：",
    "•熟练：Python / FastAPI / asyncio / PostgreSQL（pgvector、pg_trgm）/ Redis。",
)
replace_first(
    doc.paragraphs,
    "熟悉：",
    "•熟悉：Docker Compose / Nginx / SSE 流式 / Prometheus / Grafana / Tempo / Loki / "
    "Alembic / JWT 认证与权限体系 / CI/CD 流程 / 分布式锁与乐观锁。",
)

doc.save(DST)
print("saved:", DST)

# 校验：读取新文件确认关键改动
check = Document(DST)
for p in check.paragraphs:
    t = p.text.strip()
    if t and any(k in t for k in ["黎忠南", "10-12K", "本地 Locust", "熟练：", "熟悉："]):
        print("  OK:", t[:80])
