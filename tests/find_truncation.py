"""找出截断问题和低分原因"""
import openpyxl, yaml

wb = openpyxl.load_workbook(r'D:\桌面\agent_gateway\tests\ragas\civil_evaluation_report.xlsx')
ws = wb.active
h = [c.value for c in ws[1]]
id_col = 0
sc = h.index('评分(10分制)')
rc = h.index('Agent完整回复')
ac = h.index('扣分说明与综合评价')
cc = h.index('是否缓存')
ec = h.index('错误')

rows = list(ws.iter_rows(min_row=2, values_only=True))
print("=== 所有低分用例 (<8分) ===")
for r in rows:
    score = r[sc]
    if score is None or score >= 8:
        continue
    cid = r[id_col]
    resp = str(r[rc])
    cached = str(r[cc])
    err = str(r[ec]) if r[ec] else '-'
    analysis = str(r[ac])[:150]
    
    # Detect truncation: response is very short and seems cut off
    is_truncated = len(resp) < 100 and ('查询' in resp[:50] or '查一下' in resp[:50])
    
    notes = []
    if is_truncated:
        notes.append('⚠️ 疑似截断')
    if err and err != '-':
        notes.append(f'错误={err[:30]}')
    
    print(f'\n{cid}: {score}/10 len={len(resp)} {"|".join(notes)}')
    print(f'  缓存={cached} 错误={err}')
    print(f'  响应: {resp[:120]}')
    print(f'  分析: {analysis}')
