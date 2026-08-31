"""查看低分用例的详细回复"""
import openpyxl

wb = openpyxl.load_workbook(r'D:\桌面\agent_gateway\tests\evaluation_report.xlsx')
ws = wb.active
header = [cell.value for cell in ws[1]]
id_col = 0
resp_col = header.index('Agent完整回复')
tool_col = header.index('工具调用上下文')

targets = ['CC019', 'CC035', 'CC041', 'CC057', 'CC060', 'CC046', 'CC001']

for target in targets:
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[id_col] == target:
            sep = "=" * 60
            print(f'{sep}')
            print(f'{target} - Response:')
            resp = str(row[resp_col]) if row[resp_col] else '(empty)'
            print(resp[:500])
            print()
            print(f'Tool context:')
            tc = str(row[tool_col]) if row[tool_col] else '(none)'
            print(tc[:300])
            print()
            break
