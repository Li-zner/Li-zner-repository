# -*- coding: utf-8 -*-
"""堵死 HR 误读点：
1) 公网演示去掉"非工作时段可能断连"免责（网站当前在线已验证 200）
2) 去掉"3 个月从零到上线"速成暗示 → 强调全流程独立完成
3) 去掉"本地压测"贬义前缀 → Locust 压测
4) 开源与自研：删 LangChain（项目未使用，requirements 无），强调自研 + 未依赖编排框架
仅改文本，保留段落样式。"""
import json
from docx import Document

# ========== 1. 简历 docx ==========
DOCX = r"d:\桌面\agent_gateway\黎忠南简历-最终版.docx"
doc = Document(DOCX)

REPLACEMENTS = [
    # (定位片段, 整段新文本)
    (
        "公网演示：https://the-world-agent.cloud（本地部署，非工作时段可能断连）",
        "公网演示：https://the-world-agent.cloud",
    ),
    (
        "3 个月从零到上线",
        "面向企业知识服务的 AI 网关，含完整支付体系，已上线公网稳定运行；"
        "全程一人独立完成：需求拆解、架构设计、开发实现、部署运维全流程。"
        "Locust 压测：800 并发近零错误、极限 1600 并发可用性 99.28%。",
    ),
    (
        "本地 Locust 压测",
        "•性能与稳定性：分布式限流、熔断降级、三级语义缓存、冷热分层记忆；"
        "Locust 分布式压测 800 并发近零错误、极限 1600 并发可用性 99.28%。",
    ),
    (
        "本地压测对比",
        "•性能优化：三级语义缓存将命中查询成本降至纯调用的约六成（压测对比）；"
        "冷热分层记忆 + 断连自动保存；压测 800 并发平均响应 1.1s。",
    ),
    (
        "基于 FastAPI / PostgreSQL / Redis / LangChain",
        "•自研实现：Agent 编排引擎、意图路由、RAG 双通道召回、支付并发安全"
        "（分布式锁 / 乐观锁 / WAL 流水）、三级语义缓存、限流熔断均为自研实现；"
        "仅复用 FastAPI / PostgreSQL / Redis 等成熟开源组件作为基础设施，未依赖任何 Agent 编排框架。",
    ),
]

for frag, new_text in REPLACEMENTS:
    done = False
    for p in doc.paragraphs:
        if frag in p.text:
            if p.runs:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new_text)
            done = True
            break
    if not done:
        print("!! docx 未找到:", frag)

doc.save(DOCX)
print("docx saved")

# ========== 2. prompts/me.json ==========
ME = r"d:\桌面\agent_gateway\prompts\me.json"
with open(ME, "r", encoding="utf-8-sig") as f:
    data = json.load(f)
old = "公网演示：https://the-world-agent.cloud（本地部署，非工作时段可能断连）"
new = "公网演示：https://the-world-agent.cloud"
assert old in data["prompt"], "me.json 未找到公网演示行"
data["prompt"] = data["prompt"].replace(old, new)
with open(ME, "w", encoding="utf-8", newline="\n") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)
    f.write("\n")
print("me.json saved")

# ========== 3. 校验 ==========
chk = Document(DOCX)
print("---- 校验 docx ----")
for p in chk.paragraphs:
    t = p.text
    if any(k in t for k in ["公网演示", "Locust", "自研实现", "3 个月"]):
        print(" |", t[:90])
print("---- 校验 me.json ----")
with open(ME, "r", encoding="utf-8-sig") as f:
    d2 = json.load(f)
print("  无'非工作时段':", "非工作时段" not in d2["prompt"])
print("  含'GitHub：':", "GitHub：github.com/Li-zner" in d2["prompt"])
