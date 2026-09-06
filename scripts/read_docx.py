import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)
print(f"===== {path} =====")
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        print(f"[{p.style.name}] {t}")
for i, table in enumerate(doc.tables):
    print(f"--- TABLE {i+1} ---")
    for row in table.rows:
        print(" | ".join(c.text.strip() for c in row.cells))
