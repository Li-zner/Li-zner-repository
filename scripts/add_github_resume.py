# -*- coding: utf-8 -*-
"""在最终版简历里加入 GitHub 个人主页地址。
只改含"源码仓库"的一行，保留其余内容与格式。"""
from docx import Document

SRC = r"d:\桌面\agent_gateway\黎忠南简历-最终版.docx"

doc = Document(SRC)

NEW = "GitHub：github.com/Li-zner ｜ 源码仓库：github.com/Li-zner/Li-zner-repository"

found = False
for p in doc.paragraphs:
    if "源码仓库" in p.text:
        if p.runs:
            p.runs[0].text = NEW
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(NEW)
        found = True
        break

if not found:
    raise SystemExit("未找到含'源码仓库'的段落")

doc.save(SRC)
print("saved:", SRC)

# 校验
chk = Document(SRC)
for p in chk.paragraphs:
    if "GitHub" in p.text:
        print("OK:", p.text)
